# -*- coding: utf-8 -*-
"""
core.py 핵심 로직에 대한 최소 단위 테스트.

실행 방법:
    pytest

화면(tkinter)을 열지 않고, 함수 하나하나에 "이런 입력을 주면 이런 결과가 나와야 한다"를
코드로 확인한다. 기능을 추가하다가 기존 동작이 실수로 깨지면 여기서 바로 알 수 있다.
"""

import openpyxl

from core import (
    normalize,
    build_header_map,
    parse_column_spec,
    resolve_column_plan,
    iter_matched_rows,
    copy_cell,
    make_unique_sheet_name,
    run_conversion,
    RawToken,
    ColumnPlan,
)


# ------------------------------------------------------------------
# normalize
# ------------------------------------------------------------------

def test_normalize_공백_제거():
    assert normalize("  SITE  ") == "SITE"


def test_normalize_none은_빈문자열():
    assert normalize(None) == ""


def test_normalize_숫자도_문자열로():
    assert normalize(123) == "123"


# ------------------------------------------------------------------
# build_header_map
# ------------------------------------------------------------------

def test_build_header_map_헤더행에서_이름과_열번호_추출():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["No.", "SITE", "유니트명"])  # 1행

    header_map = build_header_map(ws, header_row=1)

    assert header_map == {"No.": 1, "SITE": 2, "유니트명": 3}


def test_build_header_map_빈_셀은_제외():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["No.", None, "유니트명"])

    header_map = build_header_map(ws, header_row=1)

    assert header_map == {"No.": 1, "유니트명": 3}


# ------------------------------------------------------------------
# parse_column_spec (문자열 파싱만, 시트 조회 없음)
# ------------------------------------------------------------------

def test_parse_column_spec_기본_토큰_하위호환():
    tokens = parse_column_spec("SITE,1,A")
    assert tokens == [
        RawToken(None, None, "SITE"),
        RawToken(None, None, "1"),
        RawToken(None, None, "A"),
    ]


def test_parse_column_spec_출력이름_지정():
    tokens = parse_column_spec("SITE=지역")
    assert tokens == [RawToken("SITE", None, "지역")]


def test_parse_column_spec_시트까지_지정():
    tokens = parse_column_spec("SITE=관리표!지역")
    assert tokens == [RawToken("SITE", "관리표", "지역")]


def test_parse_column_spec_빈_토큰은_스페이서():
    tokens = parse_column_spec("1,,3")
    assert tokens == [
        RawToken(None, None, "1"),
        RawToken("", None, None),
        RawToken(None, None, "3"),
    ]


def test_parse_column_spec_이름만_있고_값_없으면_이름있는_빈열():
    tokens = parse_column_spec("비고=")
    assert tokens == [RawToken("비고", None, None)]


# ------------------------------------------------------------------
# resolve_column_plan
# ------------------------------------------------------------------

def _two_sheet_workbook():
    """RMA 시트(No., SITE, SYSTEM)와 관리표 시트(No., 지역)를 가진 워크북."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "RMA"
    ws1.append(["No.", "SITE", "SYSTEM"])
    ws1.append([101, "서울", "A"])
    ws1.append([102, "부산", "B"])

    ws2 = wb.create_sheet("관리표")
    ws2.append(["No.", "지역"])
    ws2.append([102, "대구"])
    ws2.append([101, "인천"])
    return wb


def test_resolve_column_plan_기본_토큰_하위호환():
    wb = _two_sheet_workbook()
    plan, missing = resolve_column_plan(wb, "RMA", 1, "No.,SITE", log=lambda m: None)

    assert plan == [ColumnPlan("No.", "RMA", 1), ColumnPlan("SITE", "RMA", 2)]
    assert missing == []


def test_resolve_column_plan_출력이름_직접_지정():
    wb = _two_sheet_workbook()
    plan, missing = resolve_column_plan(wb, "RMA", 1, "MY_SITE=SITE", log=lambda m: None)

    assert plan == [ColumnPlan("MY_SITE", "RMA", 2)]
    assert missing == []


def test_resolve_column_plan_다른_시트_참조():
    wb = _two_sheet_workbook()
    plan, missing = resolve_column_plan(wb, "RMA", 1, "SITE=관리표!지역", log=lambda m: None)

    assert plan == [ColumnPlan("SITE", "관리표", 2)]
    assert missing == []


def test_resolve_column_plan_스페이서_열_포함():
    wb = _two_sheet_workbook()
    plan, missing = resolve_column_plan(wb, "RMA", 1, "No.,,SITE", log=lambda m: None)

    assert plan == [
        ColumnPlan("No.", "RMA", 1),
        ColumnPlan("", None, None),
        ColumnPlan("SITE", "RMA", 2),
    ]


def test_resolve_column_plan_없는_열_없는_시트는_missing():
    wb = _two_sheet_workbook()
    plan, missing = resolve_column_plan(
        wb, "RMA", 1, "X=존재안함,Y=없는시트!아무거나", log=lambda m: None
    )

    assert plan == []
    assert len(missing) == 2


def test_resolve_column_plan_비어있으면_기본_열_목록_사용():
    wb = _two_sheet_workbook()
    plan, missing = resolve_column_plan(wb, "RMA", 1, "", log=lambda m: None)

    assert ColumnPlan("No.", "RMA", 1) in plan
    assert ColumnPlan("SITE", "RMA", 2) in plan
    assert "계약년" in missing  # RMA 시트에 없는 기본 열


def _diff_header_row_workbook():
    """RMA는 2행이 헤더, SYSTEM_예시는 1행이 헤더인(서로 다른) 워크북."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "RMA"
    ws1.append(["타이틀 줄"])
    ws1.append(["No.", "SITE"])
    ws1.append([1, "서울"])

    ws2 = wb.create_sheet("SYSTEM_예시")
    ws2.append(["No.", "삼성SN"])
    ws2.append([1, "SN001"])
    return wb


def test_resolve_column_plan_시트별_헤더_행을_따로_지정하면_찾음():
    wb = _diff_header_row_workbook()
    plan, missing = resolve_column_plan(
        wb, "RMA", 2, "SITE,삼성SN=SYSTEM_예시!삼성SN", log=lambda m: None,
        sheet_header_rows={"SYSTEM_예시": 1},
    )

    assert plan == [
        ColumnPlan("SITE", "RMA", 2),
        ColumnPlan("삼성SN", "SYSTEM_예시", 2),
    ]
    assert missing == []


def test_resolve_column_plan_시트별_헤더_행_안_주면_기준_시트_행을_그대로_써서_못_찾음():
    """이 테스트는 sheet_header_rows 없이 호출하면 왜 실패하는지(=원래 버그) 보여주기 위한 회귀 테스트."""
    wb = _diff_header_row_workbook()
    plan, missing = resolve_column_plan(
        wb, "RMA", 2, "삼성SN=SYSTEM_예시!삼성SN", log=lambda m: None,
    )

    assert plan == []
    assert missing  # SYSTEM_예시의 2행(데이터 행)에서 찾으려다 실패


# ------------------------------------------------------------------
# iter_matched_rows (다중 시트 결합의 핵심 - 행 순서가 달라도 키로 짝을 맞추는지)
# ------------------------------------------------------------------

def _read(wb, loc):
    sheet, r, c = loc
    return wb[sheet].cell(row=r, column=c).value if c is not None else None


def test_iter_matched_rows_행_순서가_달라도_키로_짝을_맞춤():
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "RMA"
    ws1.append(["No.", "장비명"])
    ws1.append([103, "SK하이닉스"])
    ws1.append([101, "삼성전자"])
    ws1.append([102, "LG전자"])

    ws2 = wb.create_sheet("관리표")
    ws2.append(["No.", "지사"])
    ws2.append([101, "서울지사"])
    ws2.append([103, "대구지사"])
    ws2.append([102, "부산지사"])

    plan = [
        ColumnPlan("No.", "RMA", 1),
        ColumnPlan("장비명", "RMA", 2),
        ColumnPlan("지사", "관리표", 2),
    ]

    results = [
        [_read(wb, loc) for loc in locations]
        for locations in iter_matched_rows(wb, "RMA", 1, plan, "No.", lambda m: None)
    ]

    assert results == [
        [103, "SK하이닉스", "대구지사"],
        [101, "삼성전자", "서울지사"],
        [102, "LG전자", "부산지사"],
    ]


def test_iter_matched_rows_키가_안_맞으면_빈값():
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "RMA"
    ws1.append(["No.", "장비명"])
    ws1.append([999, "미확인장비"])

    ws2 = wb.create_sheet("관리표")
    ws2.append(["No.", "지사"])
    ws2.append([101, "서울지사"])

    plan = [ColumnPlan("No.", "RMA", 1), ColumnPlan("지사", "관리표", 2)]
    logs = []

    results = [
        [_read(wb, loc) for loc in locations]
        for locations in iter_matched_rows(wb, "RMA", 1, plan, "No.", logs.append)
    ]

    assert results == [[999, None]]
    assert any("키 값이 맞지 않아" in m for m in logs)


def test_iter_matched_rows_스페이서는_빈_위치():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "RMA"
    ws.append(["No."])
    ws.append([1])

    plan = [ColumnPlan("No.", "RMA", 1), ColumnPlan("", None, None)]
    results = list(iter_matched_rows(wb, "RMA", 1, plan, "", lambda m: None))

    assert results == [[("RMA", 2, 1), (None, None, None)]]


def test_iter_matched_rows_키가_없으면_순서대로_매칭():
    """행 맞춤 기준 열을 안 주면, 값으로 짝짓지 않고 '몇 번째 데이터 행인지'로 그대로 짝짓는다."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "RMA"
    ws1.append(["No.", "장비명"])
    ws1.append([103, "SK하이닉스"])
    ws1.append([101, "삼성전자"])
    ws1.append([102, "LG전자"])

    ws2 = wb.create_sheet("SYSTEM_예시")
    ws2.append(["No.", "계약년"])
    ws2.append([101, "2024"])
    ws2.append([103, "2022"])
    ws2.append([102, "2023"])

    plan = [
        ColumnPlan("No.", "RMA", 1),
        ColumnPlan("장비명", "RMA", 2),
        ColumnPlan("계약년", "SYSTEM_예시", 2),
    ]

    results = [
        [_read(wb, loc) for loc in locations]
        for locations in iter_matched_rows(wb, "RMA", 1, plan, "", lambda m: None)
    ]

    # No. 값과 무관하게, RMA의 n번째 데이터 행 <-> SYSTEM_예시의 n번째 데이터 행으로 그대로 짝지어짐
    assert results == [
        [103, "SK하이닉스", "2024"],
        [101, "삼성전자", "2022"],
        [102, "LG전자", "2023"],
    ]


def test_iter_matched_rows_순서기준_다른_시트에_행이_부족하면_빈값():
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "RMA"
    ws1.append(["No."])
    ws1.append([1])
    ws1.append([2])

    ws2 = wb.create_sheet("관리표")
    ws2.append(["값"])
    ws2.append(["첫줄"])  # 2번째 데이터 행이 없음

    plan = [ColumnPlan("No.", "RMA", 1), ColumnPlan("값", "관리표", 1)]
    logs = []

    results = [
        [_read(wb, loc) for loc in locations]
        for locations in iter_matched_rows(wb, "RMA", 1, plan, "", logs.append)
    ]

    assert results == [[1, "첫줄"], [2, None]]
    assert any("행이 부족해" in m for m in logs)


# ------------------------------------------------------------------
# copy_cell
# ------------------------------------------------------------------

def test_copy_cell_값만_복사():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "hello"
    src = ws["A1"]
    dst = ws["B1"]

    copy_cell(src, dst, copy_style=False, copy_formula=False)

    assert dst.value == "hello"


def test_copy_cell_수식복사_옵션_켜면_수식_그대로():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "=SUM(1,2)"
    src = ws["A1"]
    dst = ws["B1"]

    copy_cell(src, dst, copy_style=False, copy_formula=True)

    assert dst.value == "=SUM(1,2)"


def test_copy_cell_수식복사_옵션_끄면_수식_문자열_그대로_값처럼_복사():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "=SUM(1,2)"
    src = ws["A1"]
    dst = ws["B1"]

    copy_cell(src, dst, copy_style=False, copy_formula=False)

    # copy_formula=False 이면 수식 문자열을 그대로 값으로 취급해서 복사한다
    assert dst.value == "=SUM(1,2)"


# ------------------------------------------------------------------
# make_unique_sheet_name
# ------------------------------------------------------------------

def test_make_unique_sheet_name_중복없으면_그대로():
    wb = openpyxl.Workbook()

    assert make_unique_sheet_name(wb, "SYSTEM") == "SYSTEM"


def test_make_unique_sheet_name_중복되면_번호_붙임():
    wb = openpyxl.Workbook()
    wb.create_sheet("SYSTEM")
    wb.create_sheet("SYSTEM_1")

    assert make_unique_sheet_name(wb, "SYSTEM") == "SYSTEM_2"


# ------------------------------------------------------------------
# run_conversion (e2e) - 기존 시트 중간에 새 열 삽입
# ------------------------------------------------------------------

def test_run_conversion_기존_시트_중간에_새_열_삽입(tmp_path):
    filepath = tmp_path / "test.xlsx"
    wb = openpyxl.Workbook()
    ws_rma = wb.active
    ws_rma.title = "RMA"
    ws_rma.append(["No.", "SITE", "SYSTEM"])
    ws_rma.append([1, "서울", "A"])
    ws_rma.append([2, "부산", "B"])

    # 대상 시트에는 이미 No., SYSTEM 열이 있고, 그 사이에 SITE가 새로 끼워져야 함
    ws_target = wb.create_sheet("SYSTEM목록")
    ws_target.append(["No.", "SYSTEM"])

    wb.save(filepath)

    sheet_name = run_conversion(
        filepath=str(filepath),
        rma_sheet_name="RMA",
        rma_header_row=1,
        column_spec_text="No.,SITE,SYSTEM",
        target_mode="existing",
        new_sheet_name="",
        existing_sheet_name="SYSTEM목록",
        copy_style=False,
        copy_formula=False,
        log=lambda m: None,
        insert_ref_column="No.",
        insert_position="after",
    )

    wb2 = openpyxl.load_workbook(filepath)
    ws2 = wb2[sheet_name]
    headers = [ws2.cell(row=1, column=c).value for c in range(1, ws2.max_column + 1)]

    assert headers == ["No.", "SITE", "SYSTEM"]  # SITE가 No.와 SYSTEM 사이에 삽입됨
    assert [ws2.cell(row=2, column=c).value for c in range(1, 4)] == [1, "서울", "A"]
    assert [ws2.cell(row=3, column=c).value for c in range(1, 4)] == [2, "부산", "B"]
